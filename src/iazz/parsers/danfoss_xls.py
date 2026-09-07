"""Parser dei polinomi Danfoss esportati da Coolselector2.

PROBLEMA
    Danfoss non pubblica un unico polinomio EN 12900: Coolselector2
    esporta un file Excel con TRE rappresentazioni della stessa macchina:
    1. foglio "Condizioni standard"    — polinomi EN 12900 (10 coeff,
       variabili Te e Tc), una riga per frequenza, unità W;
    2. foglio "Condizioni selezionate" — come sopra ma alle condizioni
       scelte dall'utente in Coolselector, unità kW, e in più la portata
       massica M [kg/h];
    3. foglio "Condizioni selezionate, velocità inclusa" — SOLO per i
       modelli VSD: un unico polinomio a 30 coefficienti in TRE variabili
       (Te, Tc, S) che copre tutto il range di velocità.

CONCETTO
    Il formato "Extended" a 30 coefficienti è un cubo completo in Te e Tc
    moltiplicato per (1, S, S²):

        Y = C0 + C1·Te + C2·Tc + C3·Te² + C4·Tc²
            + C5·Te·Tc·S² + C6·Te²·Tc·S² + C7·Te·Tc²·S²
            + C8·Te·Tc·S  + C9·Te²·Tc·S  + C10·Te·Tc²·S
            + C11·Te·Tc + C12·Te²·Tc + C13·Te·Tc² + C14·Te³ + C15·Tc³
            + C16·S + C17·Te·S + C18·Tc·S + C19·Te²·S + C20·Tc²·S
            + C21·Te³·S + C22·Tc³·S
            + C23·S² + C24·Te·S² + C25·Tc·S² + C26·Te²·S² + C27·Tc²·S²
            + C28·Te³·S² + C29·Tc³·S²

    con Te, Tc in °C e S in GIRI AL SECONDO (rev/s = rpm/60).
    L'unità di S non è dichiarata nel file: l'abbiamo determinata
    numericamente confrontando il polinomio 3-variabili con le righe
    EN 12900 per-frequenza — con S in rev/s coincidono entro lo 0,3 %,
    con S in rpm i risultati divergono di ordini di grandezza.

SCELTE
    - Il parser legge i file .xlsx CONVERTITI (LibreOffice: `soffice
      --headless --convert-to xlsx *.xls`). Gli .xls originali Danfoss
      hanno un header OLE fuori standard che xlrd rifiuta come corrotto.
    - Dal foglio 2 ("Condizioni selezionate") prendiamo i polinomi
      per-frequenza: è l'unico con la portata M, che serve per la
      cross-validation col ciclo CoolProp.
    - I blocchi di colonne vengono individuati dalle etichette di
      riga 15 ("Q [kW]", "P [kW]", ...), non da indici fissi: così il
      parser regge sia il layout a 4 grandezze sia quello a 5.
    - L'export testato è in italiano (nomi fogli inclusi). Un export
      in altra lingua richiederà di estendere `_SHEET_SELECTED` e
      `_SHEET_SPEED` — decisione rimandata a quando servirà.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from pydantic import BaseModel

# Nomi (prefissi) dei fogli nell'export italiano di Coolselector2.
# Excel tronca i nomi foglio a 31 caratteri, quindi si confronta col prefisso.
_SHEET_SELECTED = "Condizioni selezionate"
_SHEET_SPEED = "Condizioni selezionate, velocit"

# Etichette dei blocchi (riga 15 del foglio 2, riga 22 del foglio 3)
# → (chiave interna, fattore di conversione all'unità interna).
# Unità interne: W per le potenze (coerenza con lo schema Bitzer),
# kg/h per la portata, A e °C così come sono.
_BLOCK_KEYS: dict[str, tuple[str, float]] = {
    "Q [kW]": ("Q_evap_W", 1000.0),
    "P [kW]": ("P_input_W", 1000.0),
    "Q [W]": ("Q_evap_W", 1.0),
    "P [W]": ("P_input_W", 1.0),
    "I [A]": ("current_A", 1.0),
    "M [kg/h]": ("m_dot_kg_h", 1.0),
    "T_disc [°C]": ("T_discharge_C", 1.0),
}

# Nel formato a 30 coefficienti le grandezze restano in unità "native"
# (kW, A, kg/h, °C): il valutatore le converte a valle. Qui solo i nomi.
_EXT30_KEYS: dict[str, str] = {
    "Q [kW]": "Q_evap_kW",
    "P [kW]": "P_input_kW",
    "I [A]": "current_A",
    "M [kg/h]": "m_dot_kg_h",
    "T_disc [°C]": "T_discharge_C",
}

_N_COEFF_EN12900 = 10
_N_COEFF_EXT30 = 30
_HEADER_ROW = 16       # riga con "Compressor | Refrigerant | ..."
_GROUP_ROW = 15        # riga con le etichette dei blocchi "Q [kW]" ecc.
_FIRST_DATA_ROW = 17
_EXT30_HEADER_ROW = 22  # riga con le etichette nel foglio velocità
_EXT30_FIRST_ROW = 23   # riga di C0


class DanfossWorkbook(BaseModel):
    """Contenuto rilevante di un export polinomi Coolselector2.

    Attributes:
        model: sigla del compressore (es. "KS1NV0410NDCC").
        refrigerant: refrigerante del fit (es. "R1234ze(E)").
        superheat_K: surriscaldamento delle condizioni di rating.
        subcooling_K: sottoraffreddamento delle condizioni di rating.
        rpm_per_hz: rapporto giri/frequenza del motore (es. 60 → 2 poli).
        per_frequency: polinomi EN 12900 per frequenza:
            {Hz: {chiave: [10 coefficienti]}}, unità interne (W, A, kg/h, °C).
        ext30: polinomi a 30 coefficienti (Te, Tc, S) in unità native
            (kW, A, kg/h, °C), solo per i VSD; None per velocità fissa.
    """

    model: str
    refrigerant: str
    superheat_K: float
    subcooling_K: float
    rpm_per_hz: float
    per_frequency: dict[float, dict[str, list[float]]]
    ext30: dict[str, list[float]] | None = None

    @property
    def is_vsd(self) -> bool:
        """True se il file copre più frequenze (compressore VSD)."""
        return len(self.per_frequency) > 1


def _find_sheet(wb, prefix: str):
    """Restituisce il foglio il cui nome inizia con `prefix`, o None.

    Caso limite: "Condizioni selezionate" è prefisso anche del nome del
    foglio velocità ("Condizioni selezionate, velocit..."), quindi quando
    si cerca il foglio 2 bisogna scartare i match del foglio 3.
    """
    for name in wb.sheetnames:
        if name.startswith(prefix):
            if prefix == _SHEET_SELECTED and name.startswith(_SHEET_SPEED):
                continue
            return wb[name]
    return None


def _parse_selected_sheet(ws) -> tuple[dict, dict]:
    """Estrae metadati e polinomi per-frequenza dal foglio 'selezionate'.

    Returns:
        (meta, per_frequency): metadati della macchina e dizionario
        {Hz: {chiave: [10 coeff]}} in unità interne.
    """
    # 1. Mappa colonna → blocco, dalle etichette di riga 15.
    blocks: list[tuple[int, str, float]] = []  # (col_inizio, chiave, fattore)
    for c in range(1, ws.max_column + 1):
        label = ws.cell(_GROUP_ROW, c).value
        if label in _BLOCK_KEYS:
            key, scale = _BLOCK_KEYS[label]
            blocks.append((c, key, scale))
    if not blocks:
        raise ValueError(
            f"Foglio '{ws.title}': nessun blocco riconosciuto in riga "
            f"{_GROUP_ROW}. Formato Coolselector2 diverso dall'atteso?")

    # 2. Righe dati: una per frequenza.
    meta: dict = {}
    per_frequency: dict[float, dict[str, list[float]]] = {}
    for r in range(_FIRST_DATA_ROW, ws.max_row + 1):
        model = ws.cell(r, 1).value
        if not model:
            break
        rpm = float(ws.cell(r, 3).value)
        hz = float(ws.cell(r, 4).value)
        meta = {
            "model": str(model).strip(),
            "refrigerant": str(ws.cell(r, 2).value).strip(),
            "superheat_K": float(ws.cell(r, 5).value),
            "subcooling_K": float(ws.cell(r, 7).value),
            "rpm_per_hz": rpm / hz,
        }
        polys: dict[str, list[float]] = {}
        for col0, key, scale in blocks:
            coeff = [
                float(ws.cell(r, col0 + i).value) * scale
                for i in range(_N_COEFF_EN12900)
            ]
            polys[key] = coeff
        per_frequency[hz] = polys
    return meta, per_frequency


def _parse_speed_sheet(ws) -> dict[str, list[float]]:
    """Estrae i polinomi a 30 coefficienti dal foglio 'velocità inclusa'."""
    ext30: dict[str, list[float]] = {}
    for c in range(2, ws.max_column + 1):
        label = ws.cell(_EXT30_HEADER_ROW, c).value
        if label not in _EXT30_KEYS:
            continue
        coeff = [
            float(ws.cell(_EXT30_FIRST_ROW + i, c).value)
            for i in range(_N_COEFF_EXT30)
        ]
        ext30[_EXT30_KEYS[label]] = coeff
    if not ext30:
        raise ValueError(
            f"Foglio '{ws.title}': nessuna colonna riconosciuta in riga "
            f"{_EXT30_HEADER_ROW}.")
    return ext30


def parse_workbook(path: str | Path) -> DanfossWorkbook:
    """Legge un export polinomi Coolselector2 (.xlsx convertito).

    Args:
        path: file .xlsx (convertire prima gli .xls originali con
            LibreOffice, vedi docstring del modulo).

    Returns:
        DanfossWorkbook con metadati, polinomi per-frequenza e, per i
        VSD, i polinomi a 30 coefficienti.

    Raises:
        ValueError: se il layout non corrisponde al formato atteso.
    """
    path = Path(path)
    if path.suffix.lower() == ".xls":
        raise ValueError(
            f"{path.name}: gli .xls Danfoss originali non sono leggibili "
            "direttamente (header OLE fuori standard). Convertire prima: "
            "soffice --headless --convert-to xlsx <file>.xls")
    wb = openpyxl.load_workbook(path, data_only=True)

    ws_sel = _find_sheet(wb, _SHEET_SELECTED)
    if ws_sel is None:
        raise ValueError(
            f"{path.name}: foglio '{_SHEET_SELECTED}' non trovato. "
            f"Fogli presenti: {wb.sheetnames}")
    meta, per_frequency = _parse_selected_sheet(ws_sel)

    ws_speed = _find_sheet(wb, _SHEET_SPEED)
    ext30 = _parse_speed_sheet(ws_speed) if ws_speed is not None else None

    return DanfossWorkbook(**meta, per_frequency=per_frequency, ext30=ext30)


def to_map_dict(
    dwb: DanfossWorkbook,
    T_evap_range_C: tuple[float, float],
    T_cond_range_C: tuple[float, float],
    source: str = "",
) -> dict:
    """Converte un DanfossWorkbook nel dizionario JSON interno I.A.zz.

    Per i VSD emette il formato `danfoss_ext30` (un polinomio, S come
    variabile); per i fissi il classico `en12900` a 50 Hz.

    Args:
        dwb: workbook già parsato.
        T_evap_range_C: range di validità T_evap del fit [°C]. NON è
            dichiarato nel file Danfoss: va preso dall'envelope della
            macchina in Coolselector2.
        T_cond_range_C: come sopra, per T_cond.
        source: provenienza del dato (per tracciabilità).

    Returns:
        dict pronto per json.dump, caricabile da `load_compressor`.
    """
    freqs = sorted(dwb.per_frequency)
    base = {
        "_README": [
            "Mappa compressore in formato interno I.A.zz (v0).",
            "Generato da scripts/build_compressor_db.py — NON editare a",
            "mano: rigenerare dal file Danfoss originale.",
            "ATTENZIONE: T_evap_range_C / T_cond_range_C non sono nel",
            "file Danfoss; sono l'envelope indicato a mano al momento",
            "della generazione. Verificare in Coolselector2.",
        ],
        "source": source or f"Export Coolselector2, modello {dwb.model}",
        "vendor": "Danfoss",
        "model": dwb.model,
        "compressor_type": "",
        "refrigerant": dwb.refrigerant,
        "superheat_K": dwb.superheat_K,
        "subcooling_K": dwb.subcooling_K,
        "T_evap_range_C": list(T_evap_range_C),
        "T_cond_range_C": list(T_cond_range_C),
    }
    if dwb.is_vsd:
        assert dwb.ext30 is not None, "VSD senza foglio velocità?"
        base.update({
            "vsd": True,
            "poly_format": "danfoss_ext30",
            "rated_frequency_hz": 50.0,
            "frequency_min_hz": min(freqs),
            "frequency_max_hz": max(freqs),
            "frequency_scaling": "none",
            "rpm_per_hz": dwb.rpm_per_hz,
            "polynomials": {},
            "polynomials_ext30": dwb.ext30,
        })
    else:
        (hz,) = freqs
        base.update({
            "vsd": False,
            "poly_format": "en12900",
            "rated_frequency_hz": hz,
            "frequency_min_hz": hz,
            "frequency_max_hz": hz,
            "frequency_scaling": "none",
            "polynomials": dwb.per_frequency[hz],
        })
    return base
